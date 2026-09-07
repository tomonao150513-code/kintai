"""Excel (.xlsx) 出力（openpyxl）。docs/export-spec.md。"""

import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from attendance.services.formatting import format_hms_colon, format_hours_decimal

from .rows import COLUMNS, build_rows, project_summary

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E9ECEF")


def _style_header(ws):
    for cell in ws[1]:
        cell.font = _BOLD
        cell.fill = _HEADER_FILL


def _autosize(ws):
    for column_cells in ws.columns:
        width = max(
            (len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells
        )
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 8), 40)


def _bold_last_row(ws):
    for cell in ws[ws.max_row]:
        cell.font = _BOLD


def xlsx_response(entries, date_from, date_to):
    rows, total_row = build_rows(entries)
    wb = Workbook()

    detail = wb.active
    detail.title = "勤怠明細"
    detail.append(COLUMNS)
    _style_header(detail)
    for row in [*rows, total_row]:
        detail.append([row[col] for col in COLUMNS])
    _bold_last_row(detail)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    _autosize(detail)

    summary = wb.create_sheet("集計")
    summary.append(["プロジェクト", "実働秒数", "実働時間(HH:MM:SS)", "実働時間(h)", "割合(%)"])
    _style_header(summary)
    ordered, grand = project_summary(rows)
    for name, seconds in ordered:
        pct = round(seconds / grand * 100, 1) if grand else 0
        summary.append(
            [name, seconds, format_hms_colon(seconds), format_hours_decimal(seconds), pct]
        )
    summary.append(
        ["合計", grand, format_hms_colon(grand), format_hours_decimal(grand), 100 if grand else 0]
    )
    _bold_last_row(summary)
    _autosize(summary)

    stream = io.BytesIO()
    wb.save(stream)
    response = HttpResponse(stream.getvalue(), content_type=CONTENT_TYPE)
    response["Content-Disposition"] = (
        f'attachment; filename="kintai_{date_from.isoformat()}_{date_to.isoformat()}.xlsx"'
    )
    return response
