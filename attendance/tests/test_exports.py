"""CSV / Excel エクスポート（docs/export-spec.md）。"""

import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from attendance.exports.csv import csv_response
from attendance.exports.rows import COLUMNS, build_rows, completed_entries
from attendance.exports.xlsx import xlsx_response
from attendance.models import Project, Task, TimeEntry

JST = ZoneInfo("Asia/Tokyo")
pytestmark = pytest.mark.django_db


@pytest.fixture
def owner(django_user_model):
    return django_user_model.objects.create_user(username="owner", password="pw")


@pytest.fixture
def task(owner):
    project = Project.objects.create(owner=owner, name="社内システム開発")
    return Task.objects.create(project=project, name="ログイン画面実装", status=Task.Status.DOING)


@pytest.fixture
def entries(owner, task):
    TimeEntry.objects.create(
        user=owner,
        task=task,
        start_at=datetime(2026, 9, 5, 9, 0, 0, tzinfo=JST),
        end_at=datetime(2026, 9, 5, 10, 30, 24, tzinfo=JST),
        note="詳細設計まで",
        source=TimeEntry.Source.MANUAL,
    )
    TimeEntry.objects.create(
        user=owner,
        task=task,
        start_at=datetime(2026, 9, 6, 13, 0, 0, tzinfo=JST),
        end_at=datetime(2026, 9, 6, 14, 0, 0, tzinfo=JST),
    )
    # 実行中（除外されるべき）
    TimeEntry.objects.create(user=owner, task=task, start_at=datetime(2026, 9, 7, 9, 0, tzinfo=JST))
    return owner


def test_columns_order():
    assert COLUMNS == [
        "日付",
        "開始時刻",
        "終了時刻",
        "実働秒数",
        "実働時間(HH:MM:SS)",
        "実働時間(h)",
        "プロジェクト",
        "タスク",
        "タスク状態",
        "メモ",
        "入力元",
    ]


def test_build_rows_totals_and_running_excluded(entries, owner):
    qs = completed_entries(owner, date(2026, 9, 1), date(2026, 9, 30))
    rows, total_row = build_rows(qs)
    assert len(rows) == 2  # 実行中は除外
    assert rows[0]["実働秒数"] == 5424  # 1:30:24、丸めなし
    assert rows[0]["実働時間(HH:MM:SS)"] == "1:30:24"
    assert rows[0]["入力元"] == "手動"
    assert rows[1]["入力元"] == "打刻"
    assert total_row["実働秒数"] == 5424 + 3600
    assert total_row["プロジェクト"] == "合計"
    assert total_row["日付"] == ""


def test_csv_has_bom_crlf_header_and_total(entries, owner):
    qs = completed_entries(owner, date(2026, 9, 1), date(2026, 9, 30))
    resp = csv_response(qs, date(2026, 9, 1), date(2026, 9, 30))
    raw = resp.content
    assert raw.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    text = raw.decode("utf-8-sig")
    assert "\r\n" in text
    lines = text.split("\r\n")
    assert lines[0].split(",")[0] == "日付"
    assert lines[0].startswith("日付,開始時刻,終了時刻,実働秒数")
    # 最終データ行 = 合計行
    assert "合計" in lines[3]
    assert resp["Content-Type"] == "text/csv; charset=utf-8"
    assert 'filename="kintai_2026-09-01_2026-09-30.csv"' in resp["Content-Disposition"]


def test_xlsx_readback_numeric_types_and_sheets(entries, owner):
    qs = completed_entries(owner, date(2026, 9, 1), date(2026, 9, 30))
    resp = xlsx_response(qs, date(2026, 9, 1), date(2026, 9, 30))
    assert (
        resp["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="kintai_2026-09-01_2026-09-30.xlsx"' in resp["Content-Disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["勤怠明細", "集計"]
    ws = wb["勤怠明細"]
    assert [c.value for c in ws[1]] == COLUMNS
    # 2 行目 = 最初の明細。実働秒数(4列目) と 実働時間h(6列目) は数値
    assert ws.cell(row=2, column=4).value == 5424
    assert isinstance(ws.cell(row=2, column=4).value, int)
    assert isinstance(ws.cell(row=2, column=6).value, float)
    # 最終行 = 合計
    assert ws.cell(row=ws.max_row, column=4).value == 5424 + 3600
    summary = wb["集計"]
    assert summary.cell(row=1, column=1).value == "プロジェクト"
    assert summary.cell(row=summary.max_row, column=1).value == "合計"


def test_export_view_zero_rows_still_downloads(client, owner):
    client.force_login(owner)
    resp = client.get("/export/csv/?date_from=2020-01-01&date_to=2020-01-31")
    assert resp.status_code == 200
    assert resp.content.startswith(b"\xef\xbb\xbf")
    text = resp.content.decode("utf-8-sig")
    lines = [line for line in text.split("\r\n") if line]
    assert lines[0].startswith("日付,")
    assert "合計" in lines[1]  # ヘッダ + 合計行のみ


def test_export_view_invalid_range_shows_form_no_download(client, owner):
    client.force_login(owner)
    resp = client.get("/export/csv/?date_from=2026-09-30&date_to=2026-09-01")
    assert resp.status_code == 200
    assert resp["Content-Type"].startswith("text/html")
    assert "開始日は終了日以前" in resp.content.decode()


def test_export_requires_login(client):
    assert client.get("/export/").status_code == 302
